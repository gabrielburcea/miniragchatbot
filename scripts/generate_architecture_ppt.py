"""
Generates a PowerPoint deck describing the Mini RAG Chatbot system design
and architecture, based on INTERVIEW_PREP.md.

Run:
    python scripts/generate_architecture_ppt.py

Output:
    docs/architecture/Mini_RAG_Chatbot_Architecture.pptx
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.ns import qn

# ---------------------------------------------------------------------------
# Palette
# ---------------------------------------------------------------------------
NAVY = RGBColor(0x1A, 0x1E, 0x2E)
DARK_BG = RGBColor(0x11, 0x14, 0x1F)
BLUE = RGBColor(0x4A, 0x90, 0xD9)
BLUE_DARK = RGBColor(0x1A, 0x3D, 0x5C)
GREEN = RGBColor(0x4C, 0xAF, 0x50)
GREEN_DARK = RGBColor(0x1B, 0x5E, 0x20)
AMBER = RGBColor(0xFF, 0xB3, 0x00)
AMBER_DARK = RGBColor(0x8A, 0x57, 0x00)
RED = RGBColor(0xE5, 0x39, 0x35)
RED_DARK = RGBColor(0x7F, 0x00, 0x00)
PURPLE = RGBColor(0x8E, 0x6F, 0xCE)
PURPLE_DARK = RGBColor(0x3D, 0x1F, 0x6E)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GREY = RGBColor(0xB0, 0xB4, 0xC0)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

prs = Presentation()
prs.slide_width = SLIDE_W
prs.slide_height = SLIDE_H
BLANK = prs.slide_layouts[6]


def add_slide():
    slide = prs.slides.add_slide(BLANK)
    bg = slide.background
    bg.fill.solid()
    bg.fill.fore_color.rgb = DARK_BG
    return slide


def add_text(slide, left, top, width, height, text, size=18, color=WHITE,
             bold=False, italic=False, align=PP_ALIGN.LEFT, font="Calibri",
             anchor=MSO_ANCHOR.TOP, line_spacing=1.0):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        r = p.add_run()
        r.text = line
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.italic = italic
        r.font.color.rgb = color
        r.font.name = font
    return box


def add_title(slide, title, subtitle=None):
    add_text(slide, Inches(0.5), Inches(0.25), Inches(12.3), Inches(0.8),
              title, size=30, color=WHITE, bold=True)
    if subtitle:
        add_text(slide, Inches(0.5), Inches(0.85), Inches(12.3), Inches(0.5),
                  subtitle, size=16, color=BLUE, italic=True)
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(1.35),
                                   Inches(12.3), Pt(2.5))
    line.fill.solid()
    line.fill.fore_color.rgb = BLUE
    line.line.fill.background()
    return line


def add_box(slide, left, top, width, height, text, fill, line_color=None,
            text_color=WHITE, size=13, bold=True, shape=MSO_SHAPE.ROUNDED_RECTANGLE,
            font="Calibri"):
    box = slide.shapes.add_shape(shape, left, top, width, height)
    box.fill.solid()
    box.fill.fore_color.rgb = fill
    box.line.color.rgb = line_color if line_color else fill
    box.line.width = Pt(1.5)
    box.shadow.inherit = False
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = Emu(45000)
    tf.margin_right = Emu(45000)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = line
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = text_color
        r.font.name = font
    return box


def add_arrow(slide, x1, y1, x2, y2, color=GREY, width=1.5, dashed=False, label=None):
    conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    conn.line.color.rgb = color
    conn.line.width = Pt(width)
    if dashed:
        ln = conn.line._get_or_add_ln()
        d = ln.makeelement(qn('a:prstDash'), {'val': 'dash'})
        ln.append(d)
    # arrowhead
    ln = conn.line._get_or_add_ln()
    tail = ln.makeelement(qn('a:tailEnd'), {'type': 'triangle', 'w': 'med', 'len': 'med'})
    ln.append(tail)
    if label:
        midx = (x1 + x2) / 2
        midy = (y1 + y2) / 2
        add_text(slide, midx - Inches(1.1), midy - Inches(0.22), Inches(2.2), Inches(0.4),
                  label, size=10, color=GREY, align=PP_ALIGN.CENTER)
    return conn


def add_legend(slide, items, left=Inches(0.5), top=Inches(6.95)):
    x = left
    for text, color in items:
        sw = Inches(0.28)
        chip = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, top, sw, Inches(0.2))
        chip.fill.solid()
        chip.fill.fore_color.rgb = color
        chip.line.fill.background()
        chip.shadow.inherit = False
        add_text(slide, x + sw + Inches(0.08), top - Inches(0.05), Inches(2.4), Inches(0.3),
                  text, size=11, color=GREY)
        x += Inches(2.4)


# ===========================================================================
# SLIDE 1 — Title
# ===========================================================================
s = add_slide()
add_text(s, Inches(1), Inches(2.5), Inches(11.3), Inches(1.2),
          "Mini RAG Chatbot", size=48, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
add_text(s, Inches(1), Inches(3.5), Inches(11.3), Inches(0.8),
          "System Design & Architecture", size=26, color=BLUE, align=PP_ALIGN.CENTER)
add_text(s, Inches(1), Inches(4.3), Inches(11.3), Inches(0.6),
          "Authentication  •  RBAC-Secured RAG Retrieval  •  Tool-Calling LLM  •  Streaming WebSocket",
          size=15, color=GREY, align=PP_ALIGN.CENTER, italic=True)
bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(4.5), Inches(5.1), Inches(4.3), Pt(3))
bar.fill.solid(); bar.fill.fore_color.rgb = BLUE; bar.line.fill.background()

# ===========================================================================
# SLIDE 2 — High-level architecture overview (component diagram)
# ===========================================================================
s = add_slide()
add_title(s, "High-Level Architecture", "Component overview — how the subsystems connect")

top = Inches(1.8)
h = Inches(1.0)
w = Inches(2.5)
gap = Inches(0.35)

client = add_box(s, Inches(0.5), top, Inches(1.9), h, "Client\n(Browser /\nWebSocket)", BLUE, BLUE_DARK)
server = add_box(s, Inches(2.9), top, Inches(2.1), h, "FastAPI\nWebSocket Server\napp/ws/chat.py", AMBER, AMBER_DARK, text_color=RGBColor(0,0,0))
auth = add_box(s, Inches(5.5), Inches(0.8), Inches(2.0), Inches(0.8), "Auth Module\njwt_tokens.py", BLUE, BLUE_DARK)
rag = add_box(s, Inches(5.5), top, Inches(2.0), h, "RAG Retriever\nretriever.py", GREEN, GREEN_DARK)
qdrant = add_box(s, Inches(8.0), top, Inches(1.9), h, "Qdrant\nVector DB\n+ RBAC filter", GREEN, GREEN_DARK)
llm = add_box(s, Inches(5.5), Inches(2.95), Inches(2.0), h, "LLM Provider\n(Groq)\ngroq_provider.py", PURPLE, PURPLE_DARK)
tool = add_box(s, Inches(8.0), Inches(2.95), Inches(1.9), h, "Tool\nget_employee_\ncontext()", PURPLE, PURPLE_DARK)
prompt = add_box(s, Inches(10.3), top, Inches(2.3), h, "Prompt Builder\nprompts.py\n(system prompt +\nguardrails)", AMBER, AMBER_DARK, text_color=RGBColor(0,0,0))

# connections
add_arrow(s, client.left + client.width, top + h/2, server.left, top + h/2, color=BLUE, label="WS: auth + chat msgs")
add_arrow(s, server.left + server.width/2, top, auth.left + auth.width/2 - Inches(1.2), Inches(1.2), color=BLUE)
add_arrow(s, server.left + server.width, top + h/2, rag.left, top + h/2, color=GREEN, label="to_thread()")
add_arrow(s, rag.left + rag.width, top + h/2, qdrant.left, top + h/2, color=GREEN, label="search + filter")
add_arrow(s, server.left + server.width - Inches(0.3), top + h, server.left + server.width - Inches(0.3), Inches(2.95)+Inches(0.4), color=PURPLE)
add_arrow(s, server.left + server.width - Inches(0.3), Inches(3.2), llm.left, Inches(3.2), color=PURPLE, label="stream_chat()")
add_arrow(s, llm.left + llm.width, Inches(3.45), tool.left, Inches(3.45), color=PURPLE, label="tool call")
add_arrow(s, rag.left + rag.width/2, top + h, rag.left+rag.width/2, Inches(2.7), color=AMBER, dashed=True)
add_arrow(s, rag.left+rag.width/2, Inches(2.7), prompt.left+Inches(1.0), Inches(2.7), color=AMBER, dashed=True)
add_arrow(s, prompt.left+Inches(1.0), Inches(2.7), prompt.left+Inches(1.0), top+h, color=AMBER, dashed=True, label="chunks →\nprompt")

add_text(s, Inches(0.5), Inches(4.4), Inches(12.3), Inches(0.5),
          "Flow: Client authenticates once over WebSocket → server validates JWT → each question triggers RBAC-filtered\nretrieval from Qdrant → prompt assembled → Groq LLM (2-phase: tool-call then streamed answer) → streamed back to client.",
          size=13, color=GREY)

add_legend(s, [("Transport / Server", AMBER), ("Security / Client", BLUE), ("Retrieval", GREEN), ("LLM / Tools", PURPLE)])

# ===========================================================================
# SLIDE 3 — End-to-end sequence diagram
# ===========================================================================
s = add_slide()
add_title(s, "End-to-End Request Flow", "From client message to \"done\" — 10 steps")

lanes = [
    ("Client", BLUE, Inches(0.5)),
    ("Server\n(FastAPI WS)", AMBER, Inches(2.6)),
    ("Auth", BLUE, Inches(4.7)),
    ("Retriever\n+ Qdrant", GREEN, Inches(6.6)),
    ("LLM (Groq)\n+ Tool", PURPLE, Inches(9.0)),
]
lane_w = Inches(1.9)
lane_top = Inches(1.7)
for name, color, x in lanes:
    add_box(s, x, lane_top, lane_w, Inches(0.55), name, color,
            color, size=12)
    # vertical lifeline
    ln = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x + lane_w/2, lane_top+Inches(0.55), x+lane_w/2, Inches(6.9))
    ln.line.color.rgb = RGBColor(0x3A,0x3F,0x52)
    ln.line.width = Pt(1)

steps = [
    ("1-2", "auth (JWT) → verify_token() → auth_success", Inches(0.5), Inches(4.7)),
    ("3-4", "chat message → validate ChatMessage schema", Inches(0.5), Inches(2.6)),
    ("5", "retrieve_chunks() on bg thread → embed + RBAC filter → Qdrant search", Inches(2.6), Inches(6.6)),
    ("6-7", "build_system_prompt() → stream_chat(): tool-call phase then streamed phase", Inches(2.6), Inches(9.0)),
    ("8-9", "text deltas streamed → \"done\" message", Inches(9.0), Inches(0.5)),
    ("10", "on exception at any step: \"error\" msg, connection stays open", Inches(2.6), Inches(0.5)),
]
y = Inches(2.35)
row_h = Inches(0.72)
for i, (num, desc, xa, xb) in enumerate(steps):
    yy = y + i*row_h
    badge = add_box(s, Inches(0.15), yy, Inches(0.5), Inches(0.42), num, RED, RED_DARK, size=12)
    left_x = min(xa, xb) + Inches(0.95)
    right_x = max(xa, xb) + Inches(0.95)
    arrow_color = GREEN if "5" == num else (PURPLE if "6" in num else (BLUE if "8" in num or "1" in num else RED))
    add_arrow(s, left_x, yy+Inches(0.21), right_x, yy+Inches(0.21), color=arrow_color, width=2)
    add_text(s, Inches(0.75), yy - Inches(0.02), Inches(12.0), Inches(0.4), desc, size=11.5, color=WHITE)

add_legend(s, [("Auth (closes on fail)", BLUE), ("Retrieval (RBAC)", GREEN), ("LLM / Tool", PURPLE), ("Error path", RED)])

# ===========================================================================
# SLIDE 4 — RBAC security model
# ===========================================================================
s = add_slide()
add_title(s, "RBAC & Security Model", "Access control enforced INSIDE Qdrant, not after retrieval")

q = add_box(s, Inches(0.6), Inches(1.9), Inches(2.6), Inches(1.0),
            "User question +\nJWT claims:\ndepartment, level", BLUE, BLUE_DARK)
filt = add_box(s, Inches(3.7), Inches(1.9), Inches(2.9), Inches(1.0),
               "RBAC Filter\napplied INSIDE Qdrant\n(should = OR)", AMBER, AMBER_DARK, text_color=RGBColor(0,0,0))
g1 = add_box(s, Inches(7.1), Inches(1.2), Inches(2.6), Inches(0.75),
             "own department AND\naccess_level ≤ user level", GREEN, GREEN_DARK, size=11)
g2 = add_box(s, Inches(7.1), Inches(2.15), Inches(2.6), Inches(0.75),
             "OR department = hr AND\naccess_level ≤ user level", GREEN, GREEN_DARK, size=11)
result = add_box(s, Inches(10.2), Inches(1.7), Inches(2.5), Inches(1.2),
                  "Returned to app:\nONLY authorized\nchunks", GREEN, GREEN_DARK)
blocked = add_box(s, Inches(3.7), Inches(3.4), Inches(2.9), Inches(0.85),
                   "exec / finance chunks\noutside user's access", RED, RED_DARK)

add_arrow(s, q.left+q.width, Inches(2.4), filt.left, Inches(2.4), color=BLUE)
add_arrow(s, filt.left+filt.width, Inches(2.1), g1.left, Inches(1.55), color=GREEN)
add_arrow(s, filt.left+filt.width, Inches(2.55), g2.left, Inches(2.5), color=GREEN)
add_arrow(s, g1.left+g1.width, Inches(1.55), result.left, Inches(2.0), color=GREEN)
add_arrow(s, g2.left+g2.width, Inches(2.5), result.left, Inches(2.2), color=GREEN)
add_arrow(s, blocked.left+blocked.width/2, blocked.top, blocked.left+blocked.width/2, filt.top+filt.height, color=RED, dashed=True, label="never leaves Qdrant")

add_text(s, Inches(0.6), Inches(4.6), Inches(12.1), Inches(0.5),
          "Why \"should\" (OR) not \"must\" (AND)? A finance user's own-department chunks don't also need to be hr chunks —\nOR correctly expresses \"own department (capped by level) OR hr (capped by level).\"",
          size=13, color=GREY)

kp = [
    ("Enforced at the DB layer", "query_filter passed directly into Qdrant's query_points — unauthorized chunks never exist in app memory."),
    ("Not semantic-based", "Filter constrains department/level from JWT claims, regardless of what the question asks about."),
    ("Defense-in-depth guardrails", "Prompt-injection guardrails in prompts.py are a secondary layer; RBAC filter is the hard security boundary."),
]
yk = Inches(5.3)
for i,(t,d) in enumerate(kp):
    xk = Inches(0.6) + i*Inches(4.15)
    box = add_box(s, xk, yk, Inches(3.95), Inches(1.7), "", NAVY, RGBColor(0x2A,0x2F,0x45), shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    add_text(s, xk+Inches(0.15), yk+Inches(0.1), Inches(3.65), Inches(0.4), t, size=13, bold=True, color=AMBER)
    add_text(s, xk+Inches(0.15), yk+Inches(0.55), Inches(3.65), Inches(1.1), d, size=11, color=GREY)

# ===========================================================================
# SLIDE 5 — RAG pipeline (chunking → embedding → retrieval)
# ===========================================================================
s = add_slide()
add_title(s, "RAG Pipeline", "Chunking → Embedding → Similarity Search with RBAC")

items = [
    ("PDF page\ntext", BLUE, BLUE_DARK, 1.0),
    ("chunk_text()\n550 chars,\n80 overlap", AMBER, AMBER_DARK, 1.0),
    ("Chunks\n(overlapping)", PURPLE, PURPLE_DARK, 1.0),
    ("embed_texts()\nbge-small-en-v1.5", AMBER, AMBER_DARK, 1.0),
    ("384-dim\nvectors", GREEN, GREEN_DARK, 1.0),
    ("Qdrant search\n+ RBAC filter", GREEN, GREEN_DARK, 1.0),
    ("top-k relevant,\nauthorized chunks", GREEN, GREEN_DARK, 1.0),
]
x = Inches(0.4)
bw = Inches(1.65)
bh = Inches(1.1)
yb = Inches(2.2)
prev = None
for label, fill, ln, factor in items:
    b = add_box(s, x, yb, bw, bh, label, fill, ln, size=11)
    if prev is not None:
        add_arrow(s, prev.left+prev.width, yb+bh/2, x, yb+bh/2, color=GREY, width=2)
    prev = b
    x += bw + Inches(0.28)

add_text(s, Inches(0.5), Inches(3.6), Inches(12.3), Inches(0.5),
          "Why 550/80? Each policy PDF ≈ 1 page, structured as numbered clauses; 550 chars ≈ one clause.\nToo small → clause split across chunks; too large → dilutes embedding relevance with unrelated clauses.",
          size=13, color=GREY)

qa = [
    ("Pure vector search miss", "Exact clause lookups (e.g. \"section 4.2\") can rank worse than topical matches — a lexical method (BM25) would catch literal token matches vector search misses."),
    ("Measuring retrieval quality", "Build a labeled eval set (question → correct chunk); compute precision@k / recall@k, re-run after chunking/embedding/param changes to catch regressions."),
    ("RBAC vs. ranking order", "RBAC filter applied BEFORE similarity ranking inside Qdrant — an unauthorized chunk with higher raw score is excluded from the candidate pool entirely, never just filtered out after."),
]
yq = Inches(4.3)
for i,(t,d) in enumerate(qa):
    xk = Inches(0.5) + i*Inches(4.15)
    add_box(s, xk, yq, Inches(3.95), Inches(2.6), "", NAVY, RGBColor(0x2A,0x2F,0x45))
    add_text(s, xk+Inches(0.15), yq+Inches(0.12), Inches(3.65), Inches(0.5), t, size=13, bold=True, color=AMBER)
    add_text(s, xk+Inches(0.15), yq+Inches(0.65), Inches(3.65), Inches(1.85), d, size=11, color=GREY)

# ===========================================================================
# SLIDE 6 — Async/concurrency correctness
# ===========================================================================
s = add_slide()
add_title(s, "Async & Concurrency Correctness", "Parallel tool calls + non-blocking retrieval")

add_text(s, Inches(0.5), Inches(1.7), Inches(6), Inches(0.4), "❌ Serial await (~3s total) — NOT used", size=15, bold=True, color=RED)
serial_items = ["profile()\n1s", "manager()\n1s", "team()\n1s"]
xs = Inches(0.5)
for it in serial_items:
    add_box(s, xs, Inches(2.15), Inches(1.7), Inches(0.85), it, RED, RED_DARK, size=11)
    xs += Inches(1.9)
    if xs < Inches(5.5):
        pass
for i in range(2):
    add_arrow(s, Inches(0.5)+Inches(1.7)+i*Inches(1.9), Inches(2.57), Inches(0.5)+Inches(1.9)+i*Inches(1.9), Inches(2.57), color=RED, width=2)

add_text(s, Inches(7.0), Inches(1.7), Inches(6), Inches(0.4), "✅ asyncio.gather (~1s total) — actually used", size=15, bold=True, color=GREEN)
xg = Inches(7.0)
for it in ["profile()\n1s", "manager()\n1s", "team()\n1s"]:
    add_box(s, xg, Inches(2.15), Inches(1.7), Inches(0.85), it, GREEN, GREEN_DARK, size=11)
    xg += Inches(1.9)

merged = add_box(s, Inches(9.4), Inches(3.35), Inches(2.9), Inches(0.85),
                  "merged dict: name, grade,\nmanager, team_size, team_name", BLUE, BLUE_DARK, size=11)
for i in range(3):
    add_arrow(s, Inches(7.0)+Inches(0.85)+i*Inches(1.9), Inches(3.0), Inches(9.4)+Inches(1.45), Inches(3.35), color=GREEN, width=1.2)

add_text(s, Inches(0.5), Inches(4.5), Inches(12.3), Inches(0.6),
          "Key point: asyncio.gather only parallelizes cooperative, await-yielding code (asyncio.sleep). A blocking call like\ntime.sleep(1) inside async def would still serialize execution on the single event loop thread.",
          size=13, color=GREY)

box2 = add_box(s, Inches(0.5), Inches(5.3), Inches(12.1), Inches(1.7), "", NAVY, RGBColor(0x2A,0x2F,0x45))
add_text(s, Inches(0.75), Inches(5.4), Inches(11.6), Inches(0.4),
          "Why retrieve_chunks() needed asyncio.to_thread()", size=14, bold=True, color=AMBER)
add_text(s, Inches(0.75), Inches(5.85), Inches(11.6), Inches(1.1),
          "embed_query() runs sentence-transformers .encode() synchronously on CPU, and Qdrant's query_points() is a "
          "blocking HTTP call — neither yields to the event loop. Called directly in async def ws_chat, either would "
          "freeze every other connected user's WebSocket. asyncio.to_thread() offloads the whole synchronous function "
          "to a worker thread, keeping the event loop free to serve other connections concurrently.",
          size=12, color=GREY)

# ===========================================================================
# SLIDE 7 — LLM integration & tool-calling (2-phase)
# ===========================================================================
s = add_slide()
add_title(s, "LLM Integration & Tool-Calling", "Two-phase call: non-streamed tool decision → streamed answer")

m = add_box(s, Inches(0.4), Inches(2.0), Inches(2.1), Inches(0.9), "messages:\nsystem prompt +\nuser question", BLUE, BLUE_DARK, size=11)
p1 = add_box(s, Inches(2.9), Inches(2.0), Inches(2.3), Inches(0.9), "Phase 1:\nnon-streamed call\n(decide tool use)", AMBER, AMBER_DARK, size=11, text_color=RGBColor(0,0,0))
dec = add_box(s, Inches(5.6), Inches(2.0), Inches(1.8), Inches(0.9), "tool_calls\npresent?", PURPLE, PURPLE_DARK, size=11, shape=MSO_SHAPE.DIAMOND)
tool2 = add_box(s, Inches(5.35), Inches(3.4), Inches(2.3), Inches(0.85), "tool_executor()\nget_employee_context()\n(up to 3 rounds)", PURPLE, PURPLE_DARK, size=10)
p2 = add_box(s, Inches(7.9), Inches(2.0), Inches(2.3), Inches(0.9), "Phase 2:\nstreamed call\n(stream=True)", GREEN, GREEN_DARK, size=11)
out = add_box(s, Inches(10.7), Inches(2.0), Inches(2.2), Inches(0.9), "text deltas\nyielded to\nWebSocket client", GREEN, GREEN_DARK, size=11)
err = add_box(s, Inches(5.35), Inches(4.5), Inches(2.3), Inches(0.85), "RuntimeError:\nexceeded max\ntool-call rounds", RED, RED_DARK, size=11)

add_arrow(s, m.left+m.width, Inches(2.45), p1.left, Inches(2.45), color=BLUE)
add_arrow(s, p1.left+p1.width, Inches(2.45), dec.left, Inches(2.45), color=AMBER)
add_arrow(s, dec.left+dec.width, Inches(2.3), p2.left, Inches(2.3), color=GREEN, label="no more tools")
add_arrow(s, p2.left+p2.width, Inches(2.45), out.left, Inches(2.45), color=GREEN)
add_arrow(s, dec.left+dec.width/2, dec.top+dec.height, tool2.left+tool2.width/2, tool2.top, color=PURPLE, label="yes (≤3 rounds)")
add_arrow(s, tool2.left+tool2.width/2, tool2.top, m.left+m.width/2+Inches(1.5), Inches(2.9), color=PURPLE, dashed=True, label="append result → loop")
add_arrow(s, dec.left+dec.width/2, dec.top+dec.height, err.left+err.width/2, err.top, color=RED, dashed=True, label="exceeded 3 rounds")

qa2 = [
    ("Why not tool-call + stream together?", "Streamed tool-call arguments arrive as partial JSON fragments across chunks — can't execute a half-formed tool call. Two-phase: complete tool call first, then stream final text only."),
    ("Model swap risk", "openai/gpt-oss-20b chosen since llama-3.3-70b-versatile 404'd (Enterprise-only). GROQ_MODEL is config-driven, so swapping models is a one-line .env change, no code changes."),
]
yq2 = Inches(5.6)
for i,(t,d) in enumerate(qa2):
    xk = Inches(0.4) + i*Inches(6.15)
    add_box(s, xk, yq2, Inches(5.95), Inches(1.6), "", NAVY, RGBColor(0x2A,0x2F,0x45))
    add_text(s, xk+Inches(0.15), yq2+Inches(0.08), Inches(5.65), Inches(0.4), t, size=13, bold=True, color=AMBER)
    add_text(s, xk+Inches(0.15), yq2+Inches(0.5), Inches(5.65), Inches(1.05), d, size=11, color=GREY)

# ===========================================================================
# SLIDE 8 — Error handling & resilience
# ===========================================================================
s = add_slide()
add_title(s, "Error Handling & Resilience", "Recoverable errors vs. connection-ending failures")

e = add_box(s, Inches(0.5), Inches(2.2), Inches(1.9), Inches(0.9), "Something\ngoes wrong", BLUE, BLUE_DARK, size=12)
k = add_box(s, Inches(2.8), Inches(2.2), Inches(2.0), Inches(0.9), "What kind\nof failure?", PURPLE, PURPLE_DARK, size=12, shape=MSO_SHAPE.DIAMOND)

outcomes = [
    ("invalid / expired /\nmissing JWT", "auth_failed +\nCLOSE connection", RED, RED_DARK, Inches(1.4)),
    ("malformed chat\nmessage", "error message +\ncontinue loop", AMBER, AMBER_DARK, Inches(2.5)),
    ("Qdrant\nunreachable", "error message +\nconnection stays OPEN", AMBER, AMBER_DARK, Inches(3.6)),
    ("LLM / tool\nfailure", "error message +\nconnection stays OPEN", AMBER, AMBER_DARK, Inches(4.7)),
    ("client\ndisconnects", "WebSocketDisconnect\ncaught, exits quietly", GREEN, GREEN_DARK, Inches(5.8)),
]
for cond, res, fill, ln, yy in outcomes:
    c = add_box(s, Inches(5.4), yy, Inches(2.6), Inches(0.75), cond, NAVY, RGBColor(0x2A,0x2F,0x45), size=10.5)
    r = add_box(s, Inches(8.5), yy, Inches(2.6), Inches(0.75), res, fill, ln, size=10.5,
                text_color=RGBColor(0,0,0) if fill==AMBER else WHITE)
    add_arrow(s, k.left+k.width, Inches(2.65), c.left, yy+Inches(0.37), color=GREY)
    add_arrow(s, c.left+c.width, yy+Inches(0.37), r.left, yy+Inches(0.37), color=(RED if fill==RED else (GREEN if fill==GREEN else AMBER)))

add_arrow(s, e.left+e.width, Inches(2.65), k.left, Inches(2.65), color=BLUE)

add_text(s, Inches(0.5), Inches(6.7), Inches(12.3), Inches(0.5),
          "Line drawn at trust: bad auth = no reliable identity/RBAC context → close connection. Bad message/downstream\nfailure = connection + auth still valid → send error, keep the session alive for the next question.",
          size=12, color=GREY)

# ===========================================================================
# SLIDE 9 — Testing strategy
# ===========================================================================
s = add_slide()
add_title(s, "Testing Strategy", "Fast unit tests + live integration tests")

fast_hdr = add_text(s, Inches(0.5), Inches(1.8), Inches(5.8), Inches(0.4),
                     "Fast, self-contained (no external deps)", size=15, bold=True, color=GREEN)
fast_tests = ["test_chunking.py", "test_embeddings.py", "test_auth.py", "test_tools_parallel.py"]
yft = Inches(2.3)
for t in fast_tests:
    add_box(s, Inches(0.5), yft, Inches(5.6), Inches(0.55), t, GREEN, GREEN_DARK, size=12)
    yft += Inches(0.68)

live_hdr = add_text(s, Inches(7.0), Inches(1.8), Inches(5.8), Inches(0.4),
                     "Live infrastructure (real Qdrant / Groq)", size=15, bold=True, color=AMBER)
live_tests = ["test_vectorstore.py", "test_retriever_rbac.py (top_k=50 workaround)", "test_ws_protocol.py (real Groq API call)"]
ylt = Inches(2.3)
for t in live_tests:
    add_box(s, Inches(7.0), ylt, Inches(5.6), Inches(0.55), t, AMBER, AMBER_DARK, size=12, text_color=RGBColor(0,0,0))
    ylt += Inches(0.68)

result = add_box(s, Inches(3.6), Inches(4.55), Inches(6.1), Inches(0.75),
                  "26 tests, all passing — ~15s total", BLUE, BLUE_DARK, size=15)
add_arrow(s, Inches(3.3), yft, result.left, Inches(4.9), color=GREEN)
add_arrow(s, Inches(9.8), ylt, result.left+result.width, Inches(4.9), color=AMBER)

add_text(s, Inches(0.5), Inches(5.6), Inches(12.3), Inches(1.3),
          "Tradeoff: live tests prove real RBAC/protocol correctness against actual infra, but are slower, dependent on "
          "external services being up, and slightly non-deterministic. RBAC tests use top_k=50 as a documented, pragmatic "
          "workaround to avoid being crowded out by real ingested data — the honest fix is an isolated test collection "
          "per test run.",
          size=12.5, color=GREY)

# ===========================================================================
# SLIDE 10 — Tradeoffs & judgment
# ===========================================================================
s = add_slide()
add_title(s, "Tradeoffs & Judgment", "Decisions made under the assignment's constraints")

ask = add_box(s, Inches(0.4), Inches(1.9), Inches(2.5), Inches(0.9), "Assignment suggests:\nllama-3.3-70b-versatile", BLUE, BLUE_DARK, size=11)
chk = add_box(s, Inches(3.2), Inches(1.9), Inches(2.0), Inches(0.9), "Available on\nfree-tier key?", PURPLE, PURPLE_DARK, size=11, shape=MSO_SHAPE.DIAMOND)
opt1 = add_box(s, Inches(5.6), Inches(1.4), Inches(2.9), Inches(0.85), "Option A: switch Groq model\n(openai/gpt-oss-20b)", GREEN, GREEN_DARK, size=10.5)
opt2 = add_box(s, Inches(5.6), Inches(2.5), Inches(2.9), Inches(0.85), "Option B: switch provider\nentirely (e.g. Ollama)", AMBER, AMBER_DARK, size=10.5, text_color=RGBColor(0,0,0))
r1 = add_box(s, Inches(8.9), Inches(1.4), Inches(3.4), Inches(0.85), "Same architecture, same tool-\ncalling pattern, 1-line config change", GREEN, GREEN_DARK, size=10)
r2 = add_box(s, Inches(8.9), Inches(2.5), Inches(3.4), Inches(0.85), "Different API shape, more risk,\nlikely slower (local CPU)", RED, RED_DARK, size=10)
final = add_box(s, Inches(9.8), Inches(3.6), Inches(2.5), Inches(0.7), "Chosen: Option A", BLUE, BLUE_DARK, size=13)

add_arrow(s, ask.left+ask.width, Inches(2.35), chk.left, Inches(2.35), color=BLUE)
add_arrow(s, chk.left+chk.width, Inches(2.1), opt1.left, Inches(1.82), color=RED, label="404 - Enterprise only")
add_arrow(s, chk.left+chk.width, Inches(2.6), opt2.left, Inches(2.92), color=RED, label="404 - Enterprise only")
add_arrow(s, opt1.left+opt1.width, Inches(1.82), r1.left, Inches(1.82), color=GREEN)
add_arrow(s, opt2.left+opt2.width, Inches(2.92), r2.left, Inches(2.92), color=AMBER)
add_arrow(s, r1.left+r1.width/2, r1.top+r1.height, final.left+final.width/2, final.top, color=GREEN)

points = [
    ("Blocking-call fix & error handling > stretch goals", "Both map to explicit rubric requirements. Stretch goals (reranking, hybrid search, memory) are explicitly optional — required-item correctness is the higher-leverage investment."),
    ("Highest-leverage next step: reranking + hybrid search", "BM25 + vector reranking directly improves answer correctness — the core value proposition — more than conversation memory, which improves convenience across turns rather than single-answer quality."),
]
yp = Inches(4.6)
for i,(t,d) in enumerate(points):
    xk = Inches(0.4) + i*Inches(6.15)
    add_box(s, xk, yp, Inches(5.95), Inches(2.3), "", NAVY, RGBColor(0x2A,0x2F,0x45))
    add_text(s, xk+Inches(0.15), yp+Inches(0.12), Inches(5.65), Inches(0.6), t, size=13, bold=True, color=AMBER)
    add_text(s, xk+Inches(0.15), yp+Inches(0.75), Inches(5.65), Inches(1.45), d, size=11.5, color=GREY)

# ===========================================================================
out_dir = "/Users/gabrielburcea/gen_ai_project/ai-chatbot-challenge-rfybza/miniragchatbot/docs/architecture"
os.makedirs(out_dir, exist_ok=True)
out_path = os.path.join(out_dir, "Mini_RAG_Chatbot_Architecture.pptx")
prs.save(out_path)
print(f"Saved: {out_path}")
